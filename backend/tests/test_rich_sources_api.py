import io

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from backend.tests.conftest import make_client
from backend.tests.support.http import create_subject, upload_source, wait_for_operation
from backend.tests.support.media import image_bytes


def pdf_bytes():
    output = io.BytesIO()
    document = canvas.Canvas(output)
    document.drawString(72, 720, "Derivative measures the rate of change.")
    document.save()
    return output.getvalue()


def docx_bytes():
    output = io.BytesIO()
    document = Document()
    document.add_heading("Cell biology", level=1)
    document.add_paragraph("Mitochondria produce cellular energy.")
    document.save(output)
    return output.getvalue()


def pptx_bytes():
    output = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text = "Photosynthesis converts light energy into chemical energy."
    presentation.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "content", "query", "location_kind"),
    [
        ("notes.pdf", pdf_bytes(), "Derivative", "page"),
        ("notes.docx", docx_bytes(), "Mitochondria", "section"),
        ("notes.pptx", pptx_bytes(), "Photosynthesis", "slide"),
    ],
)
def test_rich_document_anchors_are_retrievable(tmp_path, filename, content, query, location_kind):
    client, app = make_client(tmp_path)
    with client:
        subject_id = create_subject(client, "资料测试")
        accepted = upload_source(client, subject_id, filename, content)
        operation = wait_for_operation(client, accepted.json()["operation"]["id"])
        assert operation["status"] == "succeeded"
        version_id = operation["result"]["id"]

        anchors = app.state.source_library.retrieve(query, [version_id])
        assert len(anchors) == 1
        response = client.get(f"/api/source-versions/{version_id}/anchors/{anchors[0]['id']}")
        assert response.status_code == 200
        assert response.json()["location"]["kind"] == location_kind


def test_image_source_exposes_original_asset(tmp_path):
    client, _ = make_client(tmp_path)
    content = image_bytes()
    with client:
        subject_id = create_subject(client, "图片测试")
        accepted = upload_source(client, subject_id, "diagram.png", content, "image/png")
        operation = wait_for_operation(client, accepted.json()["operation"]["id"])
        version_id = operation["result"]["id"]
        version = client.get(f"/api/source-versions/{version_id}").json()

        assert version["anchor_count"] == 1
        assert version["assets"][0]["width"] == 32
        asset_id = version["assets"][0]["id"]
        downloaded = client.get(f"/api/source-versions/{version_id}/assets/{asset_id}")
        assert downloaded.headers["content-type"] == "image/png"
        assert downloaded.content == content
