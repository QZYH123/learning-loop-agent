"""Parsers that turn supported source files into neutral anchors and assets."""
from __future__ import annotations

import io
import mimetypes
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


def parse_source(media_kind: str, content: bytes) -> dict:
    if media_kind in {"markdown", "text"}:
        return _parse_text(content, markdown=media_kind == "markdown")
    if media_kind == "pdf":
        return _parse_pdf(content)
    if media_kind == "docx":
        return _parse_docx(content)
    if media_kind == "pptx":
        return _parse_pptx(content)
    if media_kind == "image":
        return _parse_image(content)
    raise ValueError(f"unsupported source kind: {media_kind}")


def _parse_text(content: bytes, *, markdown: bool) -> dict:
    text = content.decode("utf-8-sig")
    chunks = [chunk.strip() for chunk in re.split(r"\n\s*\n", text) if chunk.strip()]
    section_path: list[str] = []
    anchors = []
    for index, chunk in enumerate(chunks):
        heading = re.match(r"^(#{1,6})\s+(.+)$", chunk.splitlines()[0]) if markdown else None
        if heading:
            level = len(heading.group(1))
            section_path = [*section_path[: level - 1], heading.group(2).strip()]
        label = " / ".join(section_path) if section_path else f"段落 {index + 1}"
        anchors.append({
            "location": {
                "kind": "section" if section_path else "paragraph",
                "label": label,
                "section_path": section_path,
                "page": None,
                "slide": None,
                "block_index": index,
                "asset_index": None,
            },
            "content": [{"type": "markdown", "text": chunk}],
        })
    return {"anchors": anchors, "assets": []}


def _parse_pdf(content: bytes) -> dict:
    reader = PdfReader(io.BytesIO(content))
    anchors = []
    assets = []
    for page_number, page in enumerate(reader.pages, start=1):
        blocks = []
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append({"type": "markdown", "text": text})
        for image in page.images:
            asset_index = len(assets)
            mime_type = mimetypes.guess_type(image.name)[0] or "application/octet-stream"
            width, height = _image_size(image.data)
            assets.append({
                "kind": "image",
                "mime_type": mime_type,
                "width": width,
                "height": height,
                "alt": f"第 {page_number} 页图片",
                "data": image.data,
            })
            blocks.append({"type": "image", "asset_index": asset_index, "alt": f"第 {page_number} 页图片", "caption": None})
        if blocks:
            anchors.append({
                "location": {
                    "kind": "page",
                    "label": f"第 {page_number} 页",
                    "section_path": [],
                    "page": page_number,
                    "slide": None,
                    "block_index": page_number - 1,
                    "asset_index": None,
                },
                "content": blocks,
            })
    return {"anchors": anchors, "assets": assets}


def _parse_docx(content: bytes) -> dict:
    document = Document(io.BytesIO(content))
    anchors = []
    assets = []
    section_path: list[str] = []
    block_index = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            if item.style and item.style.name.startswith("Heading"):
                level_text = item.style.name.removeprefix("Heading").strip()
                level = int(level_text) if level_text.isdigit() else 1
                section_path = [*section_path[: level - 1], text]
            anchors.append({
                "location": {
                    "kind": "section" if section_path else "paragraph",
                    "label": " / ".join(section_path) if section_path else f"段落 {block_index + 1}",
                    "section_path": section_path,
                    "page": None,
                    "slide": None,
                    "block_index": block_index,
                    "asset_index": None,
                },
                "content": [{"type": "markdown", "text": text}],
            })
            block_index += 1
        elif isinstance(item, Table):
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            columns = rows[0] if rows else []
            anchors.append({
                "location": {
                    "kind": "table",
                    "label": f"表格 {block_index + 1}",
                    "section_path": section_path,
                    "page": None,
                    "slide": None,
                    "block_index": block_index,
                    "asset_index": None,
                },
                "content": [{"type": "table", "columns": columns, "rows": rows[1:]}],
            })
            block_index += 1

    for relation in document.part.rels.values():
        if not relation.reltype.endswith("/image"):
            continue
        data = relation.target_part.blob
        mime_type = relation.target_part.content_type
        width, height = _image_size(data)
        asset_index = len(assets)
        assets.append({
            "kind": "image",
            "mime_type": mime_type,
            "width": width,
            "height": height,
            "alt": "文档图片",
            "data": data,
        })
        anchors.append({
            "location": {
                "kind": "image",
                "label": f"图片 {asset_index + 1}",
                "section_path": [],
                "page": None,
                "slide": None,
                "block_index": block_index,
                "asset_index": asset_index,
            },
            "content": [{"type": "image", "asset_index": asset_index, "alt": "文档图片", "caption": None}],
        })
        block_index += 1
    return {"anchors": anchors, "assets": assets}


def _parse_pptx(content: bytes) -> dict:
    presentation = Presentation(io.BytesIO(content))
    anchors = []
    assets = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                asset_index = len(assets)
                width, height = _image_size(image.blob)
                assets.append({
                    "kind": "image",
                    "mime_type": image.content_type,
                    "width": width,
                    "height": height,
                    "alt": f"第 {slide_number} 张幻灯片图片",
                    "data": image.blob,
                })
                blocks.append({
                    "type": "image",
                    "asset_index": asset_index,
                    "alt": f"第 {slide_number} 张幻灯片图片",
                    "caption": None,
                })
            elif getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                blocks.append({"type": "table", "columns": rows[0] if rows else [], "rows": rows[1:]})
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    blocks.append({"type": "markdown", "text": text})
        if blocks:
            anchors.append({
                "location": {
                    "kind": "slide",
                    "label": f"第 {slide_number} 张幻灯片",
                    "section_path": [],
                    "page": None,
                    "slide": slide_number,
                    "block_index": slide_number - 1,
                    "asset_index": None,
                },
                "content": blocks,
            })
    return {"anchors": anchors, "assets": assets}


def _parse_image(content: bytes) -> dict:
    with Image.open(io.BytesIO(content)) as image:
        mime_type = Image.MIME.get(image.format, "application/octet-stream")
        width, height = image.size
    asset = {
        "kind": "image",
        "mime_type": mime_type,
        "width": width,
        "height": height,
        "alt": "图片资料",
        "data": content,
    }
    return {
        "assets": [asset],
        "anchors": [{
            "location": {
                "kind": "image",
                "label": "图片",
                "section_path": [],
                "page": None,
                "slide": None,
                "block_index": 0,
                "asset_index": 0,
            },
            "content": [{"type": "image", "asset_index": 0, "alt": "图片资料", "caption": None}],
        }],
    }


def _image_size(content: bytes) -> tuple[int | None, int | None]:
    try:
        with Image.open(io.BytesIO(content)) as image:
            return image.size
    except OSError:
        return None, None
